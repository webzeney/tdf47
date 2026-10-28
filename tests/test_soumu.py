"""fetch_soumu.py / 14_ruiji.py のテスト（総務省の実ページ・実Excelの構造を再現した最小データ）。

実行: python3 tests/test_soumu.py  （pytest でも可）
"""

import importlib.util
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "src")
sys.path.insert(0, SRC)


def _load(name, rel):
    spec = importlib.util.spec_from_file_location(name, os.path.join(SRC, rel))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


fetch = _load("fetch_soumu", "ingest/fetch_soumu.py")
ruiji = _load("ruiji", "ingest/14_ruiji.py")


# 令和6年度 市町村決算カードページ（card-25.html）の構造: 1セル1都道府県、北海道のみPDF・Excelが2つ
CARD_HTML = """<html><head><meta charset="Shift_JIS"></head><body><table>
<tr><td><a href="javascript:void(0)">×</a></td></tr>
<tr><td>■北海道<a href="/main_content/001063944.pdf">1<img alt="PDF"></a> <a href="/main_content/001063945.pdf">2<img alt="PDF"></a>
<a href="/main_content/001063946.xlsx">1<img alt="EXCEL"></a> <a href="/main_content/001063947.xlsx">2<img alt="EXCEL"></a></td>
<td><a href="/main_content/001063948.pdf">■青森県<img alt="PDF"></a> <a href="/main_content/001063949.xlsx"><img alt="EXCEL"></a></td>
<td><a href="/main_content/001063966.pdf">■埼玉県<img alt="PDF"></a> <a href="/main_content/001063967.xlsx"><img alt="EXCEL"></a></td></tr>
</table></body></html>"""

SIM_HTML = """<html><body><table><tr>
<td><a href="/main_content/001063869.pdf">■北海道<img alt="PDF"></a></td>
<td><a href="/main_content/001063879.pdf">■埼玉県<img alt="PDF"></a></td>
<td><a href="/main_content/001063882.pdf">■神奈川県<img alt="PDF"></a></td>
</tr></table></body></html>"""

INDEX_HTML = """<ul>
<li><a href="/iken/zaisei/card-25.html">令和6年度市町村決算カード</a></li>
<li><a href="/iken/zaisei/card-24.html">令和5年度市町村決算カード</a></li>
<li><a href="/iken/zaisei/similarity_card-r06.html">令和6年度類似団体比較カード</a></li>
</ul>"""

BASE = "https://www.soumu.go.jp/iken/zaisei/card-25.html"


def test_card_page():
    got = fetch.parse_card_page(CARD_HTML, BASE)
    assert [g["pref"] for g in got] == ["01", "02", "11"]
    assert got[0]["pdf"] == ["https://www.soumu.go.jp/main_content/001063944.pdf",
                             "https://www.soumu.go.jp/main_content/001063945.pdf"]
    assert len(got[0]["xlsx"]) == 2
    assert got[2]["xlsx"] == ["https://www.soumu.go.jp/main_content/001063967.xlsx"]


def test_similarity_page():
    got = fetch.parse_similarity_page(SIM_HTML, BASE)
    assert got["11"].endswith("001063879.pdf")
    assert got["14"].endswith("001063882.pdf")  # 神奈川県（「奈良県」と取り違えない）
    assert "29" not in got


def test_index_and_year():
    p = fetch.parse(INDEX_HTML)
    text, url = fetch.first_link(p, r"年度市町村決算カード", "https://www.soumu.go.jp/iken/zaisei/card.html")
    assert url == "https://www.soumu.go.jp/iken/zaisei/card-25.html"
    assert fetch.fiscal_year_in(text) == "令和6年度"
    assert fetch.fiscal_year_in("令和元年度市町村決算カード") == "令和元年度"
    assert fetch.fiscal_year_in("令和１０年度") == "令和10年度"


def test_decode_sjis():
    raw = '<meta charset="Shift_JIS"><p>蕨市</p>'.encode("cp932")
    assert "蕨市" in fetch.decode_html(raw)


def _sheet(ws, pref, low4, name, typ):
    # 実ファイル（令和6年度 埼玉県・蕨市シート）と同じ位置
    ws["BZ2"], ws["CK2"], ws["CR2"], ws["DB2"] = "都道府県名", "団体名", "市町村類型", typ
    ws["BZ5"], ws["CK5"], ws["CK7"] = pref, low4, name
    ws["CR5"], ws["DB5"] = "地方交付税種地", "46062"


def test_read_workbook():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "目次"
    wb.active["E7"] = "- 目次 -"
    _sheet(wb.create_sheet("2さいたま市"), "11", "1007", "さいたま市", "政令指定都市")
    _sheet(wb.create_sheet("21蕨市"), "11", "2232", "蕨市", "Ⅱ－３")
    _sheet(wb.create_sheet("22戸田市"), 11, 2241, "戸田市", "Ⅲ－３")  # 数値で入っていても読める
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        wb.save(path)
        got = ruiji.read_workbook(path)
    assert [g["code"] for g in got] == ["111007", "112232", "112241"]
    assert [g["page"] for g in got] == [2, 3, 4]
    assert got[1]["name"] == "蕨市" and got[1]["type"] == "Ⅱ－３"
    assert got[0]["type"] == "政令指定都市"


def test_link_text_with_image_alt():
    """リンクの中に画像（alt="EXCEL"）があっても見つけられる（2026-10-28：令和6年度のページで確認）。"""
    html = ('<a href="/main_content/001044528.xlsx">全都道府県の主要財政指標<img src="x.gif" alt="EXCEL"></a>'
            '<a href="/main_content/001044529.xlsx">全市町村の主要財政指標<img alt="EXCEL" src="y"></a>'
            '<a href="/z.xlsx">市町村主要財政指標の都道府県別平均<img alt="EXCEL"></a>')
    p = fetch.parse(html)
    base = "https://www.soumu.go.jp/menu_seisaku/toukei/02zaisei07_04000135.html"
    assert fetch.first_link(p, r"^全市町村の主要財政指標(?:EXCEL|Excel|エクセル)?$", base)[1] == \
        "https://www.soumu.go.jp/main_content/001044529.xlsx"
    assert fetch.first_link(p, r"^全都道府県の主要財政指標(?:EXCEL|Excel|エクセル)?$", base)[1] == \
        "https://www.soumu.go.jp/main_content/001044528.xlsx"


if __name__ == "__main__":
    n = 0
    for k, v in list(globals().items()):
        if k.startswith("test_") and callable(v):
            v()
            n += 1
    print(f"OK {n} tests")
