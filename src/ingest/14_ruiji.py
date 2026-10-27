#!/usr/bin/env python3
"""14_ruiji.py — 類似団体区分と決算カード・類似団体比較カードを fiscal テーブルに投入（追補 #31・#32・#34）。

入力: data/soumu/manifest.json と data/soumu/kessan-card/*.xlsx（fetch_soumu.py が用意する）。

市町村決算カード Excel は「目次」シートの後に1団体1シートが並び、PDF の各ページと1対1で対応する
（シートの並び順 = PDF のページ番号。令和6年度で全47都道府県を確認済み）。
各シートの上部に「都道府県名」「団体名」「市町村類型」の見出しがあり、値の位置は次のとおり:
  - 都道府県名 … 見出しの3行下に都道府県コード2桁（例 BZ2 → BZ5 = 11）
  - 団体名     … 見出しの3行下に団体コード下4桁（例 CK2 → CK5 = 2232）、5行下に団体名（CK7 = 蕨市）
  - 市町村類型 … 同じ行の右側に類型（例 CR2 → DB2 = Ⅱ－３ ／ 政令指定都市 ／ 中核市 ／ 施行時特例市 ／ 特別区）
→ 団体コード6桁 = 都道府県コード2桁 + 下4桁（チェックデジット込み）。

投入する行（団体ごと。fiscal_year は決算年度）:
  metric='ruiji_type'      value_text=類型（決算カードの表記そのまま）
  metric='kessan_card'     value_text=都道府県別PDFのURL#page=N
  metric='similarity_card' value_text=都道府県別PDFのURL（類似団体比較カード）
冪等（PRIMARY KEY で INSERT OR REPLACE）。新しい年度は別の行として追加され、表示は最新年度を使う。

実行: python3 src/ingest/14_ruiji.py
"""

from __future__ import annotations

import datetime
import json
import os
import sys

SRC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(SRC)
sys.path.insert(0, SRC)
import db  # noqa: E402

MANIFEST = os.path.join(ROOT, "data", "soumu", "manifest.json")
NOW = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
LABELS = ("都道府県名", "団体名", "市町村類型")


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_sheet(ws) -> dict | None:
    """1シート（1団体）から {code, name, type} を読む。見出しが無いシートは None。"""
    grid = {}
    for row in ws.iter_rows(min_row=1, max_row=12):
        for c in row:
            if c.value is not None:
                grid[(c.row, c.column)] = c.value
    pos = {}
    for (r, col), v in grid.items():
        t = "".join(_s(v).split())
        if t in LABELS and t not in pos:
            pos[t] = (r, col)
    if len(pos) < 3:
        return None
    pr, pc = pos["都道府県名"]
    nr, nc = pos["団体名"]
    tr, tc = pos["市町村類型"]
    pref2 = _s(grid.get((pr + 3, pc))).zfill(2)
    low4 = _s(grid.get((nr + 3, nc))).zfill(4)
    name = _s(grid.get((nr + 5, nc)))
    typ = ""
    for col in range(tc + 1, tc + 40):
        typ = _s(grid.get((tr, col)))
        if typ:
            break
    if not (pref2.isdigit() and low4.isdigit() and typ):
        return None
    return {"code": pref2 + low4, "name": name, "type": typ}


def read_workbook(path: str) -> list[dict]:
    """シート順に読み、PDFのページ番号（ブック内の位置・1始まり）を付ける。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for i, sn in enumerate(wb.sheetnames, start=1):
        rec = read_sheet(wb[sn])
        if rec:
            rec["page"] = i
            out.append(rec)
    wb.close()
    return out


def records_from_manifest(man: dict) -> list[dict]:
    k = man["kessan_card"]
    sim = man["similarity_card"]["prefs"]
    recs = []
    for c in k["prefs"]:
        for xlsx, pdf in zip(c["xlsx_local"], c["pdf"]):
            for r in read_workbook(os.path.join(ROOT, xlsx)):
                r["card"] = f"{pdf}#page={r['page']}"
                r["sim"] = sim.get(c["pref"])
                recs.append(r)
    return recs


def write_rows(conn, recs: list[dict], fiscal_year: str, card_page: str, sim_page: str,
               sim_year: str | None = None) -> tuple[int, int]:
    """records → fiscal。entity に無い団体はスキップして件数を返す。"""
    cur = conn.cursor()
    codes = {r["code"] for r in cur.execute("SELECT code FROM entity")}
    ok = miss = 0
    for r in recs:
        if r["code"] not in codes:
            miss += 1
            print(f"  entity未登録: {r['code']} {r.get('name', '')}", file=sys.stderr)
            continue
        rows = [("ruiji_type", r["type"], fiscal_year, card_page),
                ("kessan_card", r["card"], fiscal_year, card_page)]
        if r.get("sim"):
            rows.append(("similarity_card", r["sim"], sim_year or fiscal_year, sim_page))
        for metric, text, fy, src in rows:
            cur.execute(
                "INSERT OR REPLACE INTO fiscal(entity_code,fiscal_year,metric,value,value_text,source_url,updated_at)"
                " VALUES(?,?,?,NULL,?,?,?)",
                (r["code"], fy, metric, text, src, NOW))
        ok += 1
    conn.commit()
    return ok, miss


def main():
    if not os.path.exists(MANIFEST):
        sys.exit("data/soumu/manifest.json がありません。先に src/ingest/fetch_soumu.py を実行してください。")
    with open(MANIFEST, encoding="utf-8") as f:
        man = json.load(f)
    recs = records_from_manifest(man)
    if len(recs) < 1700:
        sys.exit(f"読み取れた団体数が少なすぎます（{len(recs)}）。Excelの様式が変わった可能性があります。投入を中止します。")
    conn = db.connect()
    db.init_schema(conn)
    ok, miss = write_rows(conn, recs, man["kessan_card"]["fiscal_year"], man["kessan_card"]["page"],
                         man["similarity_card"]["page"], man["similarity_card"]["fiscal_year"])
    conn.close()
    types = {}
    for r in recs:
        types[r["type"]] = types.get(r["type"], 0) + 1
    print(f"類似団体区分 {man['kessan_card']['fiscal_year']}: {ok} 団体投入 / entity未登録 {miss} / 区分 {len(types)} 種")


if __name__ == "__main__":
    main()
