#!/usr/bin/env python3
"""01_entity.py — entity テーブルへ団体マスタを投入（設計書 フェーズ1）。

データ取得順序の 1〜2 に相当：
  1. 総務省 全国地方公共団体コード（data/soumu/municipality.xlsx）… 都道府県＋市区町村＋政令市行政区
  2. 総務省 一部事務組合等コード（data/soumu/union_etc.xlsx）… 後期高齢者医療広域連合47
  + 郡名は日本郵便 郵便番号データ（data/jp_post/utf_ken_all.csv）から補完（衝突回避に使用）

このスクリプトがやること:
  - かな読み → slug.py でスラッグ生成
  - path を組み立て、同一県内のスラッグ衝突を検出
  - 衝突は郡名ローマ字を前置して解決（設計書 6章）し reports/slug-collisions.csv に出力
  - entity へ upsert（何度流しても同じ結果＝冪等）

実行: python3 src/ingest/01_entity.py
"""

from __future__ import annotations

import csv
import datetime
import os
import re
import sys

import openpyxl

SRC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(SRC)
sys.path.insert(0, SRC)

import db  # noqa: E402
import slug  # noqa: E402

MUNI_XLSX = os.path.join(ROOT, "data", "soumu", "municipality.xlsx")
UNION_XLSX = os.path.join(ROOT, "data", "soumu", "union_etc.xlsx")
KEN_ALL = os.path.join(ROOT, "data", "jp_post", "utf_ken_all.csv")
COLLISION_CSV = os.path.join(ROOT, "reports", "slug-collisions.csv")

# 北方領土6村の郡名。郵便番号データに無い（郵便業務がないため）が、
# 総務省コードには現行団体として収録されている。郡名は確定した政府事実なので補完する。
HOPPO_GUN = {
    "016951": ("色丹郡", "しこたんぐん"),
    "016969": ("国後郡", "くなしりぐん"),
    "016977": ("国後郡", "くなしりぐん"),
    "016985": ("択捉郡", "えとろふぐん"),
    "016993": ("紗那郡", "しゃなぐん"),
    "017001": ("蘂取郡", "しべとろぐん"),
}

# 中核市（地方自治法。政令で指定。令和5年4月1日時点の62市）。固定リストなのでクロール不要。
CHUKAKU_NAMES = {
    "函館市", "旭川市", "青森市", "八戸市", "盛岡市", "秋田市", "山形市", "福島市", "郡山市",
    "いわき市", "宇都宮市", "前橋市", "高崎市", "川越市", "川口市", "越谷市", "柏市", "船橋市",
    "八王子市", "横須賀市", "甲府市", "松本市", "長野市", "富山市", "金沢市", "福井市", "岐阜市",
    "豊橋市", "岡崎市", "一宮市", "豊田市", "大津市", "豊中市", "吹田市", "高槻市", "枚方市",
    "八尾市", "寝屋川市", "東大阪市", "尼崎市", "西宮市", "明石市", "姫路市", "奈良市", "和歌山市",
    "鳥取市", "松江市", "倉敷市", "福山市", "呉市", "下関市", "高松市", "松山市", "高知市",
    "久留米市", "長崎市", "佐世保市", "大分市", "宮崎市", "鹿児島市", "那覇市", "水戸市",
}

NOW = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")


# ── 郡名マップ（5桁コード → (郡漢字, 郡かな)） ──────────────────────────
def load_gun_map() -> dict[str, tuple[str, str]]:
    """日本郵便データから、町村の郡名を 5桁コードで引けるようにする。"""
    gun: dict[str, tuple[str, str]] = {}
    if not os.path.exists(KEN_ALL):
        print(f"  ! 郡名ソースが無い: {KEN_ALL}（衝突解決に支障）", file=sys.stderr)
        return gun
    with open(KEN_ALL, encoding="utf-8") as f:
        for row in csv.reader(f):
            if len(row) < 8:
                continue
            code5 = row[0].strip()
            city_kana = row[4]      # ヒヤマグンエサシチョウ
            city_kanji = row[7]     # 檜山郡江差町
            if code5 in gun or "郡" not in city_kanji:
                continue
            m_kanji = re.match(r"^(.+?郡)", city_kanji)
            m_kana = re.match(r"^(.+?グン)", city_kana)
            if m_kanji and m_kana:
                gun[code5] = (m_kanji.group(1), slug.to_hiragana(m_kana.group(1)))
    return gun


def gun_for(code6: str, type_label: str, gun_map) -> tuple[str | None, str | None]:
    """町村のみ郡名を返す（市・特別区・区は郡を持たない）。"""
    if type_label not in ("町", "村"):
        return None, None
    if code6 in HOPPO_GUN:
        return HOPPO_GUN[code6]
    return gun_map.get(code6[:5], (None, None))


def load_zip_map() -> dict[str, str]:
    """日本郵便データから、市区町村の代表郵便番号7桁を5桁コードで引く。

    「以下に掲載がない場合」の行を代表とし、無ければその団体の最初の行で補完する。
    """
    rep: dict[str, str] = {}
    first: dict[str, str] = {}
    if not os.path.exists(KEN_ALL):
        return rep
    with open(KEN_ALL, encoding="utf-8") as f:
        for row in csv.reader(f):
            if len(row) < 9:
                continue
            code5 = row[0].strip()
            zip7 = row[2].strip()
            if not (zip7.isdigit() and len(zip7) == 7):
                continue
            first.setdefault(code5, zip7)
            if "以下に掲載がない場合" in row[8]:
                rep[code5] = zip7
    for code5, z in first.items():
        rep.setdefault(code5, z)
    return rep


# ── 総務省：都道府県＋市区町村 ───────────────────────────────────────
def load_prefectures_and_munis():
    wb = openpyxl.load_workbook(MUNI_XLSX, read_only=True, data_only=True)
    ws = wb["R6.1.1現在の団体"]
    prefs = {}   # 2桁プレフィックス → dict(code,name,kana,slug)
    munis = []   # 市区町村（特別区含む）
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = str(r[0]).strip()
        pname, mname, pkana, mkana = r[1], r[2], r[3], r[4]
        if mname is None:  # 都道府県行
            prefs[code[:2]] = {
                "code": code, "name": pname,
                "kana": slug.to_hiragana(pkana),
                "slug": slug.pref_slug(pname, pkana),
            }
        else:
            s, _suf, tl = slug.muni_slug(mname, mkana)
            munis.append({
                "code": code, "pref2": code[:2], "name": mname,
                "kana": slug.to_hiragana(mkana), "base_slug": s, "type_label": tl,
            })
    wb.close()
    return prefs, munis


# ── 総務省：政令指定都市の行政区 ─────────────────────────────────────
def load_wards():
    """sheet2 から行政区を抽出。親政令市の code / 名 / かなを前置除去してスラッグ化。

    returns: (wards, seirei_city_codes)  seirei=政令市のcode集合（type_label補正用）
    """
    wb = openpyxl.load_workbook(MUNI_XLSX, read_only=True, data_only=True)
    ws = wb["R6.1.1政令指定都市"]
    wards = []
    seirei = set()
    cur = None  # 現在の親政令市
    for r in ws.iter_rows(min_row=2, values_only=True):
        code = str(r[0]).strip()
        name, kana = r[2], slug.to_hiragana(r[4])
        if name.endswith("市"):           # 政令市のヘッダ行
            cur = {"code": code, "name": name, "kana": kana, "pref2": code[:2]}
            seirei.add(code)
            continue
        if not name.endswith("区") or cur is None:
            continue
        # 親市名・親かなを前置除去 → 区の部分だけ
        ward_name = name[len(cur["name"]):]
        if not kana.startswith(cur["kana"]):
            raise ValueError(f"区かなが親市かなで始まらない: {name} / {kana} / 親{cur['kana']}")
        ward_kana = kana[len(cur["kana"]):]
        s, _suf, tl = slug.muni_slug(ward_name, ward_kana)
        wards.append({
            "code": code, "pref2": code[:2], "parent_code": cur["code"],
            "name": name, "kana": kana, "base_slug": s, "type_label": "区",
        })
    wb.close()
    return wards, seirei


# ── 総務省：後期高齢者医療広域連合（1県1団体に集約） ──────────────────
def load_unions(prefs):
    """55件ある会計分割エントリを都道府県ごとに1件へ集約して47件にする。"""
    pref_by_name = {p["name"]: p for p in prefs.values()}
    pref_names = sorted(pref_by_name, key=len, reverse=True)  # 長い名から照合
    wb = openpyxl.load_workbook(UNION_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    picked: dict[str, dict] = {}  # pref2 → union dict
    for r in ws.iter_rows(min_row=4, values_only=True):
        code, name, furi, addr = r[1], r[2], r[3], r[6]
        if not name or "後期高齢者医療広域連合" not in str(name):
            continue
        pref = next((p for p in pref_names if str(addr).startswith(p)), None)
        if pref is None:
            raise ValueError(f"広域連合の所在地から都道府県を判定できない: {name} / {addr}")
        p = pref_by_name[pref]
        has_paren = "（" in str(name)
        cand = {
            "code": str(code).strip(), "pref2": p["code"][:2],
            "pref_code": p["code"], "no_paren": not has_paren,
        }
        cur = picked.get(p["code"][:2])
        # 括弧なし（本体）を優先。無ければ最小コード。
        if cur is None or (cand["no_paren"] and not cur["no_paren"]) \
                or (cand["no_paren"] == cur["no_paren"] and cand["code"] < cur["code"]):
            picked[p["code"][:2]] = cand
    wb.close()
    return list(picked.values())


# ── 衝突検出＆解決 ──────────────────────────────────────────────────
def resolve_collisions(entities, prefs):
    """同一県内で slug が重複する市区町村を、郡名前置で解決する。

    entities: path 未確定の市区町村（kind=city）。dict を直接更新して final slug を入れる。
    returns: collision_rows（レポート用）
    """
    from collections import defaultdict
    groups = defaultdict(list)
    for e in entities:
        groups[(e["pref2"], e["base_slug"])].append(e)

    rows = []
    for (pref2, base), members in groups.items():
        if len(members) == 1:
            members[0]["slug"] = base
            continue
        pslug = prefs[pref2]["slug"]
        for e in members:
            if e["gun_kana"]:
                g = slug.gun_slug(e["gun_kanji"], e["gun_kana"])
                e["slug"] = f"{g}-{base}"
            else:
                e["slug"] = base  # 郡名が無く解決不能 → 後段の一意性チェックで検出
            rows.append({
                "pref": prefs[pref2]["name"], "pref_slug": pslug,
                "base_slug": base, "code": e["code"], "name": e["name"],
                "kana": e["kana"], "gun": e["gun_kanji"] or "",
                "resolved_slug": e["slug"],
                "resolved_path": f"/{pslug}/{e['slug']}/",
            })
    return rows


def write_collision_report(rows):
    os.makedirs(os.path.dirname(COLLISION_CSV), exist_ok=True)
    cols = ["pref", "pref_slug", "base_slug", "code", "name", "kana",
            "gun", "resolved_slug", "resolved_path"]
    with open(COLLISION_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in sorted(rows, key=lambda x: (x["pref_slug"], x["base_slug"], x["code"])):
            w.writerow(r)


# ── entity への冪等 upsert ──────────────────────────────────────────
UPSERT_COLS = ["code", "kind", "pref_code", "parent_code", "name", "kana",
               "slug", "path", "gun", "type_label", "phase", "rep_zip"]


def upsert(conn, records):
    cur = conn.cursor()
    # 既存DBに rep_zip 列が無ければ追加（冪等マイグレーション）
    have = {r[1] for r in cur.execute("PRAGMA table_info(entity)")}
    if "rep_zip" not in have:
        cur.execute("ALTER TABLE entity ADD COLUMN rep_zip TEXT")
    existing = {r["code"]: dict(r) for r in cur.execute(
        "SELECT " + ",".join(UPSERT_COLS + ["created_at"]) + " FROM entity")}
    ins = upd = same = 0
    for rec in records:
        old = existing.get(rec["code"])
        if old is None:
            cols = UPSERT_COLS + ["created_at", "updated_at"]
            vals = [rec.get(c) for c in UPSERT_COLS] + [NOW, NOW]
            cur.execute(f"INSERT INTO entity ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)
            ins += 1
        elif any((old.get(c) or None) != (rec.get(c) or None) for c in UPSERT_COLS):
            sets = ",".join(f"{c}=?" for c in UPSERT_COLS) + ",updated_at=?"
            cur.execute(f"UPDATE entity SET {sets} WHERE code=?",
                        [rec.get(c) for c in UPSERT_COLS] + [NOW, rec["code"]])
            upd += 1
        else:
            same += 1
    conn.commit()
    return ins, upd, same


def main():
    print("== フェーズ1: 団体マスタ投入 ==")
    gun_map = load_gun_map()
    print(f"郡名マップ: {len(gun_map)} 件（日本郵便）＋ 北方領土 {len(HOPPO_GUN)} 件")
    zip_map = load_zip_map()
    print(f"代表郵便番号マップ: {len(zip_map)} 件（日本郵便）")

    prefs, munis = load_prefectures_and_munis()
    wards, seirei = load_wards()
    unions = load_unions(prefs)
    print(f"読込: 都道府県 {len(prefs)} / 市区町村 {len(munis)} / 行政区 {len(wards)} / 広域連合 {len(unions)}")

    records = []

    # 都道府県（phase v1.0）
    for p in prefs.values():
        last = p["name"][-1]
        records.append({
            "code": p["code"], "kind": "pref", "pref_code": p["code"],
            "parent_code": None, "name": p["name"], "kana": p["kana"],
            "slug": p["slug"], "path": f"/{p['slug']}/", "gun": None,
            "type_label": last if last in "都道府県" else None, "phase": "v1.0",
            "rep_zip": None,
        })

    # 政令市は市全体の郵便番号がken_allに無いので、先頭の行政区の代表番号で補完する。
    seirei_zip: dict[str, str] = {}
    for w in sorted(wards, key=lambda x: x["code"]):
        z = zip_map.get(w["code"][:5])
        if z:
            seirei_zip.setdefault(w["parent_code"], z)

    # 市区町村（特別区含む、phase v1.0）
    for m in munis:
        g_kanji, g_kana = gun_for(m["code"], m["type_label"], gun_map)
        m["gun_kanji"], m["gun_kana"] = g_kanji, g_kana
    collision_rows = resolve_collisions(munis, prefs)
    for m in munis:
        if m["code"] in seirei:
            tl = "政令指定都市"
        elif m["name"] in CHUKAKU_NAMES:
            tl = "中核市"
        else:
            tl = m["type_label"]
        pslug = prefs[m["pref2"]]["slug"]
        records.append({
            "code": m["code"], "kind": "city", "pref_code": prefs[m["pref2"]]["code"],
            "parent_code": None, "name": m["name"], "kana": m["kana"],
            "slug": m["slug"], "path": f"/{pslug}/{m['slug']}/",
            "gun": m["gun_kanji"], "type_label": tl, "phase": "v1.0",
            "rep_zip": zip_map.get(m["code"][:5]) or seirei_zip.get(m["code"]),
        })

    # 政令市の行政区（phase v1.5：マスタには入れるが v1.0 では独立ページを生成しない）
    parent_path = {m["code"]: f"/{prefs[m['pref2']]['slug']}/{m['slug']}/" for m in munis}
    for w in wards:
        w["slug"] = w["base_slug"]  # 市の下にネストするので県内衝突は起きない
        records.append({
            "code": w["code"], "kind": "ward", "pref_code": prefs[w["pref2"]]["code"],
            "parent_code": w["parent_code"], "name": w["name"], "kana": w["kana"],
            "slug": w["slug"], "path": f"{parent_path[w['parent_code']]}{w['slug']}/",
            "gun": None, "type_label": "区", "phase": "v1.5",
            "rep_zip": zip_map.get(w["code"][:5]),
        })

    # 後期高齢者医療広域連合（phase v1.0、slug は koukikourei 固定）
    for u in unions:
        pslug = prefs[u["pref2"]]["slug"]
        records.append({
            "code": u["code"], "kind": "union", "pref_code": u["pref_code"],
            "parent_code": None, "name": f"{prefs[u['pref2']]['name']}後期高齢者医療広域連合",
            "kana": "", "slug": "koukikourei",
            "path": f"/{pslug}/union/koukikourei/", "gun": None,
            "type_label": "広域連合", "phase": "v1.0", "rep_zip": None,
        })

    # path 一意性チェック（全pathが一意＝フェーズ1完了条件）
    seen = {}
    dups = []
    for rec in records:
        if rec["path"] in seen:
            dups.append((rec["path"], seen[rec["path"]], rec["code"]))
        seen[rec["path"]] = rec["code"]
    write_collision_report(collision_rows)
    print(f"衝突レポート: {COLLISION_CSV}（{len(collision_rows)} 行）")
    if dups:
        print(f"!! path 重複 {len(dups)} 件（解決できていない）:", file=sys.stderr)
        for path, a, b in dups:
            print(f"   {path}  <- {a} / {b}", file=sys.stderr)
        sys.exit(1)
    print(f"全 path 一意: OK（{len(records)} 団体）")

    conn = db.connect()
    db.init_schema(conn)
    ins, upd, same = upsert(conn, records)
    conn.close()
    print(f"投入: 新規 {ins} / 更新 {upd} / 変化なし {same}  合計 {len(records)}")


if __name__ == "__main__":
    main()
