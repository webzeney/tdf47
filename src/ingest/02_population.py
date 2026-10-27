#!/usr/bin/env python3
"""02_population.py — 人口・世帯数を entity に投入（設計書 データ取得順3 / e-Stat 住基）。

総務省 住民基本台帳人口・世帯数（市区町村別・総計）data/estat/population.xlsx を取り込む。
団体コード6桁が entity.code と一致する。人口は「計」列、世帯数は世帯列。
政令市は本体と行政区の両方が入っているが、entity 単位で値を持つので二重計上は起きない。

列: A=団体コード / B=都道府県名 / C=市区町村名 / F=人口計 / G=世帯数（0始まりで 5, 6）
実行: python3 src/ingest/02_population.py
"""

from __future__ import annotations

import datetime
import os
import sys

import openpyxl

SRC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(SRC)
sys.path.insert(0, SRC)
import db  # noqa: E402

XLSX = os.path.join(ROOT, "data", "estat", "population.xlsx")
AS_OF = "2026-01-01"  # 令和8年1月1日
NOW = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = {}
    for r in ws.iter_rows(min_row=8, values_only=True):
        code = r[0]
        if not code or not str(code).isdigit() or len(str(code)) != 6:
            continue
        pop, hh = r[5], r[6]
        rows[str(code)] = (int(pop) if pop not in (None, "") else None,
                           int(hh) if hh not in (None, "") else None)
    wb.close()

    conn = db.connect()
    db.init_schema(conn)
    cur = conn.cursor()
    existing = {r["code"]: dict(r) for r in cur.execute(
        "SELECT code,population,households,pop_as_of FROM entity")}
    upd = same = miss = 0
    for code, (pop, hh) in rows.items():
        old = existing.get(code)
        if old is None:
            miss += 1
            continue
        if old["population"] == pop and old["households"] == hh and old["pop_as_of"] == AS_OF:
            same += 1
            continue
        cur.execute("UPDATE entity SET population=?, households=?, pop_as_of=?, updated_at=? WHERE code=?",
                    (pop, hh, AS_OF, NOW, code))
        upd += 1
    conn.commit()
    filled = cur.execute("SELECT count(*) FROM entity WHERE population IS NOT NULL").fetchone()[0]
    conn.close()
    print(f"人口・世帯数: 更新 {upd} / 変化なし {same} / entity未登録 {miss}")
    print(f"人口が入っている entity: {filled}")


if __name__ == "__main__":
    main()
