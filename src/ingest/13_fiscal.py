#!/usr/bin/env python3
"""13_fiscal.py — 主要財政指標を fiscal テーブルに投入（追補 #33 / 総務省 主要財政指標一覧）。

取り込むファイル（data/soumu/）:
  - shihyo-shichoson.xlsx  総務省「全市町村の主要財政指標」(main_content/001044529.xlsx)
      R2ヘッダ: 団体コード | 都道府県名 | 団体名 | 財政力指数 | 経常収支比率 | 実質公債費比率 | 将来負担比率 | ラスパイレス指数
      → 団体コード6桁が entity.code と一致（1行1団体）。
  - shihyo-todofuken.xlsx  総務省「全都道府県の主要財政指標」(main_content/001044528.xlsx)
      R2ヘッダ: 都道府県名 | 財政力指数 | ...（団体コード無し）→ 都道府県名で entity に突き合わせ。

基準年度(fiscal_year)はファイル先頭A1の「令和N年度」を抽出。数値でない欄（将来負担比率の「-」等）は投入しない。
冪等（PRIMARY KEY で INSERT OR REPLACE）。

実行: python3 src/ingest/13_fiscal.py
"""

from __future__ import annotations

import datetime
import os
import re
import sys

import openpyxl

SRC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.dirname(SRC)
sys.path.insert(0, SRC)
import db  # noqa: E402

SOUMU = os.path.join(ROOT, "data", "soumu")
NOW = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")

# 市町村ファイルの列（団体コード, 都道府県名, 団体名, 財政力, 経常収支, 実質公債費, 将来負担, ラスパイレス）
SHICHOSON = {
    "file": "shihyo-shichoson.xlsx",
    "src": "https://www.soumu.go.jp/main_content/001044529.xlsx",
    "code_col": 0, "metric_start": 3,
}
TODOFUKEN = {
    "file": "shihyo-todofuken.xlsx",
    "src": "https://www.soumu.go.jp/main_content/001044528.xlsx",
    "name_col": 0, "metric_start": 1,
}
# metric_start からの並び順（5指標）
METRICS = ["zaisei_ryoku", "keijo_shushi", "kosaihi", "shorai_futan", "laspeyres"]

Z2H = str.maketrans("０１２３４５６７８９", "0123456789")


def fiscal_year_of(ws) -> str:
    title = str(ws["A1"].value or "")
    m = re.search(r"令和[\d０-９]+年度", title)
    return m.group(0).translate(Z2H) if m else "令和6年度"


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "－", "―", "…", "n.a.", "―."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def apply_manifest():
    """fetch_soumu.py の manifest があれば、最新年度のURL・保存先を使う（無ければ上の既定値）。"""
    path = os.path.join(SOUMU, "manifest.json")
    if not os.path.exists(path):
        return
    import json
    with open(path, encoding="utf-8") as f:
        shi = json.load(f).get("shihyo") or {}
    for conf, key in ((SHICHOSON, "shichoson"), (TODOFUKEN, "todofuken")):
        if shi.get(key):
            conf["src"] = shi[key]
        if shi.get(key + "_local"):
            conf["file"] = os.path.relpath(os.path.join(ROOT, shi[key + "_local"]), SOUMU)


def main():
    apply_manifest()
    conn = db.connect()
    db.init_schema(conn)
    cur = conn.cursor()
    codes = {r["code"] for r in cur.execute("SELECT code FROM entity")}
    pref_code_by_name = {r["name"]: r["code"]
                         for r in cur.execute("SELECT code,name FROM entity WHERE kind='pref'")}

    rows = []  # (entity_code, fiscal_year, metric, value, source_url)
    miss = 0

    # --- 市町村 ---
    wb = openpyxl.load_workbook(os.path.join(SOUMU, SHICHOSON["file"]), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    fy = fiscal_year_of(ws)
    for r in ws.iter_rows(min_row=3, values_only=True):
        code = r[SHICHOSON["code_col"]]
        if code is None:
            continue
        code = str(code).strip()
        if not code.isdigit() or len(code) != 6:
            continue
        if code not in codes:
            miss += 1
            continue
        for i, metric in enumerate(METRICS):
            v = num(r[SHICHOSON["metric_start"] + i])
            if v is not None:
                rows.append((code, fy, metric, v, SHICHOSON["src"]))
    wb.close()

    # --- 都道府県（団体コード無し→名称で突き合わせ） ---
    wb = openpyxl.load_workbook(os.path.join(SOUMU, TODOFUKEN["file"]), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    fy_p = fiscal_year_of(ws)
    for r in ws.iter_rows(min_row=3, values_only=True):
        name = r[TODOFUKEN["name_col"]]
        if not name:
            continue
        code = pref_code_by_name.get(str(name).strip())
        if not code:
            miss += 1
            continue
        for i, metric in enumerate(METRICS):
            v = num(r[TODOFUKEN["metric_start"] + i])
            if v is not None:
                rows.append((code, fy_p, metric, v, TODOFUKEN["src"]))
    wb.close()

    for code, fy_, metric, v, src in rows:
        cur.execute(
            "INSERT OR REPLACE INTO fiscal(entity_code,fiscal_year,metric,value,value_text,source_url,updated_at)"
            " VALUES(?,?,?,?,NULL,?,?)",
            (code, fy_, metric, v, src, NOW))
    conn.commit()

    n_ent = cur.execute("SELECT count(DISTINCT entity_code) FROM fiscal").fetchone()[0]
    n_row = cur.execute("SELECT count(*) FROM fiscal").fetchone()[0]
    conn.close()
    print(f"主要財政指標: {n_row} 行投入 / {n_ent} 団体 / entity未登録スキップ {miss}")


if __name__ == "__main__":
    main()
